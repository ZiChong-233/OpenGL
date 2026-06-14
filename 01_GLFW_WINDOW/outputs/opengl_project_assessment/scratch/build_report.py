from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"G:\LearingofOpenGL\leraningTrajectory\01_GLFW_WINDOW\outputs\opengl_project_assessment\documents\OpenGL项目整体评估与内存释放优化方案.docx"


COLORS = {
    "blue": RGBColor(46, 116, 181),
    "dark_blue": RGBColor(31, 77, 120),
    "ink": RGBColor(20, 31, 43),
    "muted": RGBColor(89, 96, 105),
    "light_fill": "F2F4F7",
    "blue_fill": "E8EEF5",
    "risk_fill": "FCE4E4",
    "ok_fill": "E8F3EC",
    "warn_fill": "FFF4D6",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_font(run, size=None, bold=None, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_para(doc, text="", style=None, bold=False, color=None, size=None, align=None):
    p = doc.add_paragraph(style=style)
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, size=10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, size=10.5)
    return p


def add_callout(doc, title, body, fill="F4F6F9"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=10.5, bold=True, color=COLORS["dark_blue"])
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_font(r2, size=10.5, color=COLORS["ink"])
    doc.add_paragraph()
    return table


def add_table(doc, headers, rows, widths=None, header_fill="F2F4F7"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    if widths is None:
        widths = [9360 // len(headers)] * len(headers)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_width(hdr[i], widths[i])
        set_cell_margins(hdr[i])
        set_cell_shading(hdr[i], header_fill)
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_font(r, size=9.5, bold=True, color=COLORS["ink"])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, size=9.2, color=COLORS["ink"])
    doc.add_paragraph()
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, COLORS["blue"], 16, 8),
        ("Heading 2", 13, COLORS["blue"], 12, 6),
        ("Heading 3", 12, COLORS["dark_blue"], 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("OpenGL 项目评估 | 2026-05-25")
    set_font(r, size=8, color=COLORS["muted"])


def build():
    doc = Document()
    configure_document(doc)

    title = add_para(doc, "OpenGL 项目整体评估与内存释放优化方案", size=22, bold=True, color=COLORS["ink"])
    title.paragraph_format.space_after = Pt(8)
    add_para(doc, "基于当前工程源码、CMake 配置、README 记录与已有构建产物的静态分析", size=11, color=COLORS["muted"])
    add_para(doc, "项目路径：G:\\LearingofOpenGL\\leraningTrajectory\\01_GLFW_WINDOW", size=9.5, color=COLORS["muted"])
    add_para(doc, "交付日期：2026-05-25", size=9.5, color=COLORS["muted"])

    add_callout(
        doc,
        "一句话结论",
        "该项目已经具备学习型 OpenGL 小引擎的核心骨架：窗口与输入、相机控制、场景树、Mesh/Geometry/Material、Shader、Texture、Assimp 模型导入、基础光照、透明队列、Framebuffer 与 ImGui 调试入口均已出现。当前主要短板不在“能不能画”，而在资源所有权、生命周期边界、可测试性和工程化约束尚未稳定。",
        "E8EEF5",
    )

    doc.add_heading("1. 评估范围与限制", level=1)
    add_para(
        doc,
        "本报告没有修改项目 C++ 源码。分析依据包括根目录 CMakeLists.txt、main.cpp、application、glframework、wrapper、imgui、assets 目录和 README.txt 中现有记录。尝试重新构建时，当前终端环境未找到 cmake 命令，因此本文不声称“本次重新编译通过”。已有 out/build/x64-Debug 下存在 openglStudy.exe、静态库和构建日志，说明该工程此前至少完成过一次本地构建。",
    )
    add_table(
        doc,
        ["维度", "当前观察", "可信度"],
        [
            ["源码结构", "可直接读取，模块边界清晰：application、glframework、wrapper、imgui、assets、thirdParty。", "高"],
            ["运行行为", "未在本轮启动 OpenGL 窗口；只依据代码路径和已有日志判断。", "中"],
            ["构建状态", "已有构建产物，但本轮命令行找不到 cmake，无法复验。", "中"],
            ["内存释放", "可从构造/析构、new/delete、glGen/glDelete、Texture 缓存和 shutdown 顺序进行静态判断。", "高"],
        ],
        widths=[1700, 5900, 1760],
    )

    doc.add_heading("2. 项目已经完成的内容", level=1)
    add_table(
        doc,
        ["能力模块", "已完成内容", "代表文件"],
        [
            ["工程组织", "CMake 拆出 wrapper、application、glframework、imgui 子库，第三方依赖和 assets 能复制到构建目录。", "CMakeLists.txt"],
            ["窗口与上下文", "Application 封装 GLFW 初始化、OpenGL 4.6 Core Profile、GLAD 加载、事件回调和主循环更新。", "application/Application.cpp"],
            ["输入与相机", "实现 PerspectiveCamera、OrthographicCamera、TrackBallCameraControl、GameCameraControl，支持鼠标、键盘、滚轮控制。", "application/camera/*"],
            ["场景组织", "Object 支持父子树、局部变换、父矩阵组合；Scene 继承 Object，Mesh 继承 Object。", "glframework/object.cpp, scene.cpp, mesh.cpp"],
            ["几何资源", "Geometry 可创建 Box、Sphere、Plane、ScreenPlane，也支持 Assimp 顶点数据构建 VAO/VBO/EBO。", "glframework/geometry.cpp"],
            ["材质与 Shader", "MaterialType 覆盖 Phong、White、Depth、OpacityMask、Screen、Cube、PhongEnv；Renderer 可按材质选择 Shader。", "glframework/material/*, renderer.cpp"],
            ["纹理系统", "Texture 支持文件、内存、Cubemap、Framebuffer color/depth-stencil attachment，并有路径缓存。", "glframework/texture.cpp"],
            ["模型导入", "AssimpLoader 递归解析节点、变换、Mesh、Diffuse/Specular 贴图，能加载外部或嵌入纹理。", "application/assimpLoader.cpp"],
            ["渲染流程", "Renderer 每帧收集不透明/透明对象，按材质设置深度、模板、混合、面剔除状态，并执行 glDrawElements。", "glframework/renderer/renderer.cpp"],
            ["调试 UI", "集成 ImGui，已有 Clear Color 面板，主循环末尾渲染 UI。", "main.cpp"],
        ],
        widths=[1500, 5650, 2210],
    )

    doc.add_heading("3. 当前欠缺的工程能力", level=1)
    add_para(doc, "从成熟引擎的标准看，当前工程更像“课程项目逐步抽象出的渲染框架”，还没有完全进入稳定引擎层。主要欠缺集中在以下方面：")
    for item in [
        "资源所有权没有统一表达。Scene/Object/Mesh、Material/Texture、Framebuffer/Texture、Renderer/Shader 都在使用裸指针和手动 delete，但哪些指针拥有资源、哪些只是借用，规则还没有被类型系统约束。",
        "OpenGL GPU 资源类默认可拷贝。Geometry、Texture、Shader、Framebuffer 持有 GLuint 句柄，但没有禁用拷贝，也没有实现移动语义；一旦被值拷贝，就可能出现重复 glDelete 或悬空句柄。",
        "Asset 与 Render Resource 尚未分层。Texture 同时承担图片资源、GL 对象、缓存条目和纹理单元状态，材质通过 setUnit 修改共享纹理对象，容易在多材质复用时产生状态串扰。",
        "错误处理和失败路径不完整。Shader 文件读取失败后构造函数直接 return，AssimpLoader 在读取失败时已经创建的 rootNode 会泄漏，GL 调用失败缺少集中诊断报告。",
        "可测试性不足。没有自动化测试、资源泄漏计数、反复加载卸载测试、窗口 resize/FBO 重建测试，也没有构建环境的脚本化复现说明。",
        "渲染架构仍较集中。Renderer::renderObject 中材质分支较长，随着 PBR、阴影、实例化、后处理增多，会变成维护瓶颈。",
        "Application 单例生命周期不完整。Application::getInstance 使用 new 创建静态单例，但没有对应 delete；destroy 只 glfwTerminate，没有显式 glfwDestroyWindow 和状态清空。",
    ]:
        add_bullet(doc, item)

    add_table(
        doc,
        ["成熟度项", "评分", "判断"],
        [
            ["窗口/输入/基础循环", "3/5", "能跑主循环和回调，抽象基本形成；生命周期和错误恢复仍偏手动。"],
            ["渲染功能", "3/5", "具备多材质、多 Shader、光照、透明排序、FBO 入口；渲染管线还未模块化。"],
            ["资源加载", "2.5/5", "Assimp 与 stb_image 已接入；资产缓存、热更新、失败占位、格式策略不足。"],
            ["场景组织", "2.5/5", "树结构清楚；所有权、移除节点、遍历性能和组件化还未成熟。"],
            ["内存/资源释放", "2/5", "已写出大量析构释放代码，但所有权规则还不可靠。"],
            ["调试/验证", "2/5", "有 ImGui 和日志；缺少资源统计、GL debug label、自动泄漏检查。"],
            ["工程可复现", "2/5", "CMake 已有；本轮环境找不到 cmake，依赖和构建路径还需脚本化。"],
        ],
        widths=[2400, 1000, 5960],
    )

    doc.add_heading("4. 内存释放与资源生命周期专项评估", level=1)
    doc.add_heading("4.1 当前做得好的地方", level=2)
    for item in [
        "Object 析构函数递归 delete 子节点，Scene 删除时可以释放整棵对象树。",
        "Mesh 析构函数删除自己的 Geometry 和 Material，说明 Mesh 被设计为几何与材质的拥有者。",
        "Geometry 析构函数检查并删除 VAO、position/uv/normal VBO 和 EBO。",
        "Shader 构造中链接完成后删除中间 shader object，析构中删除 program。",
        "Texture 析构中删除 GL texture，stb_image 的 CPU 图片数据在上传后 stbi_image_free。",
        "Texture::clearCache 可集中释放路径缓存纹理，main.cpp 的 shutdown 在场景、渲染器释放后调用该函数。",
        "Framebuffer 析构删除 FBO 和 color/depth-stencil attachment。",
        "main.cpp 中已经有 shutdown(AppRuntime&) 统一释放运行时对象，且对指针置空。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("4.2 主要风险", level=2)
    add_table(
        doc,
        ["风险", "当前表现", "影响", "优先级"],
        [
            ["所有权语义不清", "Material 保存 Texture*，但析构不释放；Framebuffer 又拥有自己的 Texture*；Texture cache 也拥有 Texture*。", "后续扩展时容易 double delete、泄漏或悬空。", "P0"],
            ["共享纹理带状态", "Texture::setUnit 会修改 Texture 对象自身，多个材质复用同一纹理时 unit 不是调用局部状态。", "渲染结果可能受最后一次 setUnit 影响。", "P0"],
            ["GPU 资源类可拷贝", "Geometry、Texture、Shader、Framebuffer 没有禁用 copy。", "一旦被复制，两个对象会持有同一 GL id。", "P0"],
            ["失败路径泄漏", "AssimpLoader::load 先 new rootNode，再 ReadFile；失败返回 nullptr 时 rootNode 未 delete。", "模型加载失败会泄漏 CPU 对象。", "P1"],
            ["Application 单例未释放", "minstance 由 new 创建，但没有 delete；destroy 只 glfwTerminate。", "进程退出时 OS 可回收，但不是成熟引擎做法。", "P1"],
            ["调试替换材质有隐患", "setModelWhiteDebug 直接 delete mesh->mMaterial 并换新材质。", "如果未来材质拥有纹理或共享资源，可能破坏所有权。", "P1"],
            ["Context 顺序依赖隐式", "GL 对象析构必须发生在 OpenGL context 有效期间，目前靠 shutdown 在 appGI->destroy 前执行。", "顺序一旦变化，析构中的 glDelete 可能失效。", "P1"],
            ["缓存释放不可观测", "Texture cache 只有 clearCache，没有引用计数、泄漏报告、重复 key 统计。", "难以定位长期运行中的资源增长。", "P2"],
        ],
        widths=[1900, 3750, 2550, 1160],
        header_fill="E8EEF5",
    )

    add_callout(
        doc,
        "专项结论",
        "当前项目已经有“释放动作”，但还没有稳定的“生命周期模型”。成熟度问题不是缺少 glDelete，而是缺少谁拥有、谁借用、何时创建、何时销毁、失败时怎么回滚、共享资源如何引用的统一规则。",
        "FFF4D6",
    )

    doc.add_heading("5. 成熟引擎项目应该怎样做", level=1)
    add_para(doc, "成熟引擎通常不会把资源释放写成零散 delete，而会建立一套可验证的生命周期体系。对 OpenGL 项目尤其要同时管理 CPU 对象和 GPU 句柄。")
    add_table(
        doc,
        ["原则", "成熟做法", "落到本项目的含义"],
        [
            ["RAII", "C++ 对象构造获得资源，析构释放资源；异常和提前返回也不会泄漏。", "Geometry/Texture/Shader/Framebuffer 应成为不可拷贝、可移动的 RAII 包装。"],
            ["所有权显式", "独占用 unique_ptr，共享用 shared_ptr/handle，观察关系用裸指针或 weak_ptr，但必须命名清楚。", "Object children 用 unique_ptr；Material 持有 Texture handle；Renderer 队列只观察 Mesh。"],
            ["资源管理器", "纹理、Shader、Mesh、材质由 ResourceManager/AssetManager 管理，缓存有 key、引用计数和释放策略。", "Texture::mTextureCache 升级为 ResourceManager 的一部分，绑定 unit 从 Texture 对象里拆出。"],
            ["上下文安全", "GPU 资源销毁发生在 GL context 有效且渲染线程可控的阶段。", "shutdown 顺序要固定：停止渲染、销毁场景引用、flush 资源管理器、销毁 GL 对象、最后 destroy window/context。"],
            ["移动语义", "GL id 只能被一个对象拥有；转移所有权时旧对象清零。", "禁用拷贝构造/赋值，必要时实现 move constructor/move assignment。"],
            ["延迟销毁", "跨线程或帧内仍可能被 GPU 使用的资源进入 deletion queue，到安全帧点统一释放。", "后续如加入异步加载或渲染线程，需要 RenderDevice 级别的 deferred release。"],
            ["可观测性", "每类资源有 live count、峰值、debug name、来源路径和泄漏报告。", "为 Texture/Geometry/Shader/FBO 记录创建位置、GL id 和销毁日志。"],
        ],
        widths=[1500, 4100, 3760],
    )

    doc.add_heading("6. 面向当前项目的优化方案", level=1)
    add_para(doc, "建议分阶段推进，先修正所有权模型，再重构接口。这样能避免一次性大改导致渲染功能回退。")
    steps = [
        ("阶段 0：建立资源清单，不改变行为", "列出每个类的 owned/non-owned 成员；为 Material::mDiffuse 等字段标注“非拥有”；给 shutdown 顺序写注释；增加资源 live count 日志。"),
        ("阶段 1：封住最危险的复制路径", "给 Geometry、Texture、Shader、Framebuffer 删除 copy constructor 和 copy assignment；如需要容器移动，再补 move 语义。"),
        ("阶段 2：修复明显泄漏和窗口生命周期", "AssimpLoader::load 失败路径释放 rootNode 或改为验证成功后再创建；Application::destroy 增加 glfwDestroyWindow、mWindow = nullptr；考虑释放 Application 单例。"),
        ("阶段 3：拆分 Texture 资源和绑定状态", "把 Texture::bind() 改为 bind(unit) 或由 Renderer/MaterialPass 指定 unit；禁止通过共享 Texture::setUnit 改变全局状态。"),
        ("阶段 4：材质资源改为 handle/shared_ptr", "Material 不再裸持有 Texture*，改为 TextureHandle 或 std::shared_ptr<Texture>；Framebuffer attachment 明确为 unique_ptr<Texture> 或独立 RenderTexture 类型。"),
        ("阶段 5：Scene 所有权现代化", "Object::mChildren 改为 std::vector<std::unique_ptr<Object>>；addChild 返回引用或裸观察指针；Mesh 的 Geometry/Material 使用 unique_ptr。"),
        ("阶段 6：ResourceManager 与诊断工具", "统一加载 Texture/Shader/Model；提供 acquire/release、reload、clearUnused、dumpLeaks；接入 ImGui 面板展示资源数量和显存估算。"),
        ("阶段 7：压力测试", "写加载/卸载 100 次模型、重复切换材质、FBO resize、退出重进等测试，确认 CPU/GPU 资源计数回到零。"),
    ]
    for title_text, body in steps:
        p = doc.add_paragraph()
        r = p.add_run(title_text)
        set_font(r, size=10.5, bold=True, color=COLORS["dark_blue"])
        r2 = p.add_run("：" + body)
        set_font(r2, size=10.5, color=COLORS["ink"])

    doc.add_heading("7. 推荐优先级清单", level=1)
    add_table(
        doc,
        ["优先级", "建议事项", "理由"],
        [
            ["P0", "禁用 GL 资源类拷贝；明确 Texture unit 不再是共享纹理的可变状态。", "这是最容易演变成 double delete 或渲染串扰的问题。"],
            ["P0", "为 Material/Texture 制定所有权规则：缓存纹理由 ResourceManager 释放，材质只持有 handle。", "README 已记录该问题，且它会影响后续所有材质扩展。"],
            ["P1", "修复 AssimpLoader 失败泄漏、Application 窗口销毁和单例释放。", "改动小，收益明确，可作为工程化第一批提交。"],
            ["P1", "把 setModelWhiteDebug 改为 Renderer 全局覆盖材质或可恢复的调试材质栈。", "避免调试功能破坏真实材质所有权。"],
            ["P2", "建立 ResourceTracker 和 ImGui 资源面板。", "让泄漏从“猜测”变为“可观测”。"],
            ["P2", "拆分 Renderer::renderObject 中的材质分支，形成 Material::apply 或 Pass/Technique。", "为 PBR、阴影、后处理和多 pass 扩展做准备。"],
        ],
        widths=[1100, 4300, 3960],
        header_fill="E8EEF5",
    )

    doc.add_heading("8. 建议的最终状态", level=1)
    add_callout(
        doc,
        "目标架构",
        "Application 只负责窗口和循环；Renderer/RenderDevice 负责 GL 状态和安全销毁；Scene 拥有对象树；Mesh 独占 Geometry 与 Material；Material 持有资源 handle；ResourceManager 管理 Texture/Shader/Model 的缓存、引用和卸载；Debug 面板展示所有资源计数和泄漏报告。",
        "E8F3EC",
    )
    add_para(doc, "达到该状态后，项目会从“学习用框架”更接近“可持续扩展的小型引擎”：新增材质、模型、FBO、后处理或热更新时，不需要重新猜测释放责任。")

    doc.add_heading("附录：关键源码观察点", level=1)
    add_table(
        doc,
        ["文件", "观察点"],
        [
            ["main.cpp", "AppRuntime 裸指针集中管理；shutdown 手动 delete 场景、相机、光源、Renderer、Framebuffer，并调用 Texture::clearCache。"],
            ["glframework/object.cpp", "Object 析构递归删除 mChildren；addChild 防重复和多父节点。"],
            ["glframework/mesh.cpp", "Mesh 析构删除 mGeometry 和 mMaterial。"],
            ["glframework/geometry.cpp", "析构删除 VAO/VBO/EBO；构造和静态工厂负责上传顶点数据。"],
            ["glframework/texture.cpp", "有静态纹理缓存、stbi_image_free、glDeleteTextures；但 unit 是 Texture 自身状态。"],
            ["glframework/shader.cpp", "链接后删除 shader object，析构删除 program；文件读取失败路径较弱。"],
            ["glframework/material/phongMaterial.cpp", "注释已明确 mDiffuse/mSpecularMask 所有权未最终固定。"],
            ["application/assimpLoader.cpp", "加载模型并创建 Object/Mesh/Geometry/Material；读取失败时 rootNode 泄漏。"],
            ["glframework/framebuffer/framebuffer.cpp", "Framebuffer 拥有 FBO 和附件 Texture；析构删除。"],
            ["glframework/renderer/renderer.cpp", "Renderer 拥有多个 Shader；每帧临时保存 Mesh* 渲染队列，不拥有 Mesh。"],
        ],
        widths=[3000, 6360],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
